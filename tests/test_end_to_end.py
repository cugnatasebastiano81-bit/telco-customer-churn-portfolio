"""System test for the distributed model and synthetic Excel example."""

from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path
import re
import subprocess
import sys

import numpy as np
from openpyxl import load_workbook
import pandas as pd
import pytest

from predict_churn import esegui


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "predict_churn.py"
MODEL = ROOT / "pipeline_churn.joblib"
MODULE = ROOT / "churn_inference.py"
INPUT = ROOT / "examples" / "customers_sample.xlsx"
NOTEBOOK = ROOT / "telco_customer_churn_analysis.ipynb"
README = ROOT / "README.md"
README_EN = ROOT / "README_EN.md"
MODEL_CARD = ROOT / "MODEL_CARD.md"
KAGGLE = ROOT / "KAGGLE.md"

EXPECTED_HASHES = {
    MODEL: "9f808b02efd58ce2085cb5d9220e3b68abf24d5f3b66654f08d04f009c698d91",
    MODULE: "a111483858fb69cc8c782ad316a29e592cf4028ed7922927979103ba162dec40",
    INPUT: "b20e3319534ea5be57ae7f34629c7335bb81d4abf03958d5f38bc881fffb06ba",
}
EXPECTED_IDS = ["DEMO-004", "DEMO-001", "DEMO-006", "DEMO-008", "DEMO-003"]
EXPECTED_PROBABILITIES = np.array([
    0.9913397198565522,
    0.9661574233301461,
    0.8506536417210188,
    0.6202441032679623,
    0.5050068333572493,
])
EXPECTED_DECILES = np.array([1, 1, 2, 4, 4])
EXPECTED_REASON_MARKERS = {
    "DEMO-004": (
        "spesa mensile rapportata all'anzianità: sopra la media",
        "anzianità: sotto la media",
        "connessione in fibra: sì",
    ),
    "DEMO-001": (
        "spesa mensile rapportata all'anzianità: sopra la media",
        "anzianità: sotto la media",
        "connessione in fibra: sì",
    ),
    "DEMO-006": (
        "anzianità: sotto la media",
        "spesa mensile rapportata all'anzianità: sopra la media",
        "connessione in fibra: sì",
    ),
    "DEMO-008": (
        "anzianità: sotto la media",
        "spesa mensile rapportata all'anzianità: sopra la media",
        "servizio telefonico: no",
    ),
    "DEMO-003": (
        "anzianità: sotto la media",
        "familiari a carico: no",
        "contratto mensile: sì",
    ),
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def notebook_text() -> tuple[str, str]:
    notebook = json.loads(NOTEBOOK.read_text(encoding="utf-8"))
    sources = "\n".join(
        "".join(cell.get("source", [])) for cell in notebook["cells"]
    )
    outputs = []
    for cell in notebook["cells"]:
        for output in cell.get("outputs", []):
            text = output.get("text", "")
            outputs.append("".join(text) if isinstance(text, list) else text)
    return sources, "\n".join(outputs)


def markdown_section(text: str, heading: str) -> str:
    """Restituisce il contenuto di una sezione Markdown di livello 2."""
    match = re.search(
        rf"^## {re.escape(heading)}\s*$\n(.*?)(?=^## |\Z)",
        text,
        flags=re.MULTILINE | re.DOTALL,
    )
    assert match, f"sezione Markdown non trovata: {heading}"
    return match.group(1)


def test_notebook_prepara_il_path_kaggle_prima_di_importare_il_modulo():
    notebook = json.loads(NOTEBOOK.read_text(encoding="utf-8"))
    cella_setup = next(
        "".join(cell.get("source", []))
        for cell in notebook["cells"]
        if "from churn_inference import valida_dataset_sorgente" in "".join(
            cell.get("source", [])
        )
    )

    ricerca = 'Path("/kaggle/input").rglob("churn_inference.py")'
    inserimento = "sys.path.insert(0, str(moduli_churn[0].parent))"
    assert ricerca in cella_setup
    assert "if len(moduli_churn) != 1:" in cella_setup
    assert inserimento in cella_setup
    assert cella_setup.index(ricerca) < cella_setup.index(inserimento)
    assert cella_setup.index(inserimento) < cella_setup.index(
        "from churn_inference import valida_dataset_sorgente"
    )


def test_allowlist_kaggle_include_tutti_i_requirements_ricorsivi():
    kaggle = KAGGLE.read_text(encoding="utf-8")
    upload = markdown_section(kaggle, "Files to upload")
    elencati = set(re.findall(r"`([^`]+)`", upload))

    da_controllare = [
        nome for nome in elencati if nome.startswith("requirements")
    ]
    controllati = set()
    while da_controllare:
        nome = da_controllare.pop()
        if nome in controllati:
            continue
        controllati.add(nome)
        percorso = ROOT / nome
        assert percorso.is_file(), f"requirement elencato ma assente: {nome}"

        for riga in percorso.read_text(encoding="utf-8").splitlines():
            include = re.match(r"^\s*-r\s+([^#\s]+)", riga)
            if include:
                dipendenza = include.group(1)
                assert dipendenza in elencati, (
                    f"{nome} include {dipendenza}, ma il file non è "
                    "nell'allowlist Kaggle"
                )
                da_controllare.append(dipendenza)


def test_notebook_valida_il_dataset_sorgente_subito_dopo_la_lettura():
    notebook = json.loads(NOTEBOOK.read_text(encoding="utf-8"))
    caricamento = next(
        "".join(cella.get("source", []))
        for cella in notebook["cells"]
        if cella.get("cell_type") == "code"
        and "pd.read_excel(percorso_dati)" in "".join(cella.get("source", []))
    )

    assert "valida_dataset_sorgente" in caricamento
    assert caricamento.index("pd.read_excel(percorso_dati)") < caricamento.index(
        "valida_dataset_sorgente("
    )


def test_narrative_numeriche_notebook_sono_allineate_agli_output():
    sources, outputs = notebook_text()

    assert re.search(
        r"Profilo Rischio\s+13\.0\s+77\.5\s+3\s+13\s+20", outputs
    )
    assert "rank 3 per il gain, rank 20 per SHAP" in sources
    assert "appena **13 split**" in sources
    assert (
        "`Charge Per Tenure` compare in **707 split**, `Monthly Charges` in "
        "645, `Total Charges` in 517"
    ) in sources
    assert "`Monthly Charges` è rank 17 per gain medio ma rank 5 per SHAP" in sources
    assert "**0.95**, contro lo **0.72**" in sources
    assert "rank 20 in XGBoost" in sources
    assert (
        "rank 20 per SHAP in XGBoost (rank 3 per gain medio)" in sources
    )
    assert "639 predizioni positive su 1.409 osservazioni (45,4%)" in sources
    assert "79,4%" in sources
    assert "circa 3 volte la media" in sources

    obsoleti = (
        "rank 2 per il gain, rank 21 per SHAP",
        "appena **7 split**",
        "`Charge Per Tenure` compare in **709 split**",
        "rank 15 per gain medio ma rank 5 per SHAP",
        "**0.96**, contro lo **0.71**",
        "rank 21 in XGBoost",
        "Coefficiente −0.124 (penultimo in Logistica) ma rank 2 in XGBoost",
        "53% flaggato",
        "decile 1 (77%)",
        "~78% sul test",
        "oltre 3×",
    )
    for testo in obsoleti:
        assert testo not in sources


def test_claim_notebook_sono_coerenti_con_i_limiti_della_model_card():
    sources, outputs = notebook_text()
    module = MODULE.read_text(encoding="utf-8")
    presented_text = "\n".join(
        [
            sources,
            outputs,
            module,
            SCRIPT.read_text(encoding="utf-8"),
            README.read_text(encoding="utf-8"),
            MODEL_CARD.read_text(encoding="utf-8"),
        ]
    )

    assert "eventuali proxy indiretti e disparità tra sottogruppi non sono stati valutati" in sources
    assert "robustezza interna sul medesimo campione" in sources
    assert "non garantisce prestazioni operative future" in sources
    assert "non verifica il meccanismo che produce il churn" in sources
    assert "ipotesi di intervento da sottoporre a verifica" in sources
    assert "non è pronto per un uso operativo o commerciale" in sources
    assert "non misura la stabilità familiare" in sources
    assert "non stabiliscono che il volume sia gestibile" in sources
    assert "valore atteso dalla costruzione dell'esempio" in sources
    assert "scenario di costo puramente ipotetico" in sources
    assert "non dimostra che il ranking resti stabile" in sources
    assert "associazioni suggeriscono segmenti da analizzare" in sources
    assert "Conteggi descrittivi sul solo holdout" in sources
    assert "Non sono trigger operativi validati" in sources
    assert "profilo a punteggio basso" in sources
    assert "descrive la posizione rispetto alla baseline del modello" in module

    claim_non_supportati = (
        "conferma indipendente che le relazioni catturate sono reali",
        "Nessuna feature superstite fa da proxy indiretto del genere",
        "il team può fidarsi molto",
        "La funzione è pronta per il team commerciale",
        "lista pronta per il team commerciale",
        "già l'oggetto da mandare in produzione",
        "sa già su quale leva agire",
        "micro-personas azionabili",
        "questi clienti si possono lasciar stare",
        "Qui quel rischio non esiste",
        "l'unica dipendenza di codice della produzione",
        "Spiegazione più plausibile:",
        "Lo verifico con i dati di spesa",
        "segmento a rischio strutturale, da curare attivamente",
        "Il modello passa al team una lista",
        "quasi totalità dei churner reali",
        "Un trade-off vantaggioso",
        "robusto tra 2:1 e 10:1",
        'segnale "famiglia stabile"',
        "Lettura corretta:",
        "classificato correttamente come probabile churn",
        "capire **perché alcuni clienti abbandonano",
        "l'azienda può concentrare le azioni",
        "traduco i risultati in azioni concrete",
        "prima barriera al churn",
        "La finestra critica per intervenire",
        "il cliente tende a restare da solo",
        "chi paga in automatico è più stabile",
        "più servizi = più fedeltà",
        "Il cross-selling aiuta la retention",
        "Le predizioni restano quindi utilizzabili",
        "Il modello decide a chi mandare un'azione",
        "Rapporto realistico:",
        "la conclusione è robusta",
        "decili operativi",
        "# Fase 5 — Interpretazione e raccomandazioni",
        "il contratto mensile alza il rischio in modo sistematico",
        "cioè lo protegge",
        "è un segnale pessimo",
        "Per priorizzare i clienti è sufficiente che il ranking resti stabile",
        "il ranking regge, i valori assoluti no",
        "probabilità certa, sottile",
        "i tre fattori che più alzano il rischio di quel cliente",
        "La soglia operativa arriva dall'artefatto",
        "primo segmento su cui costruire una lista di contatti",
        "non hanno ancora maturato fedeltà",
        "protezione più forte",
        "i primi mesi sono critici",
        "forte effetto protettivo",
        "protezione residua modesta",
        "protezione forte",
        "vantaggio economico osservato",
        "salvataggi potenziali",
        "spesa retention 'sprecata'",
        "perdite silenziose",
        "Clienti da contattare",
        "se il budget è limitato, si parte da qui",
        "ri-addestrare subito",
        "clienti sopra soglia (da contattare)",
        "arriverebbero al team commerciale",
        "cliente fedele",
        "Lista di intervento",
        "prioritized retention list",
        "from model selection to retention actions",
        "Decili contattati",
        "Quanto lontano nella lista arrivare?",
    )
    for claim in claim_non_supportati:
        assert claim not in presented_text

    assert "dice come agire" not in module


@pytest.mark.parametrize(
    ("readme_path", "row_label"),
    (
        (README, "XGBoost, benchmark bloccato"),
        (README_EN, "XGBoost, locked benchmark"),
    ),
)
def test_metriche_xgboost_sincronizzate_tra_readme_e_notebook(
    readme_path,
    row_label,
):
    readme = readme_path.read_text(encoding="utf-8")
    _, outputs = notebook_text()

    readme_match = re.search(
        rf"^\| {re.escape(row_label)} \| "
        r"([0-9.]+) \| ([0-9.]+) \| ([0-9.]+) \| ([0-9.]+) \| "
        r"([0-9.]+) \| ([0-9.]+) \| ([0-9]+) \| ([0-9]+) \| ([0-9]+) \|$",
        readme,
        flags=re.MULTILINE,
    )
    notebook_match = re.search(
        r"^XGBoost \(controllo\)\s+"
        r"([0-9.]+)\s+([0-9.]+)\s+([0-9.]+)\s+([0-9.]+)\s+"
        r"([0-9.]+)\s+([0-9.]+)\s+([0-9]+)\s+([0-9]+)\s+([0-9]+)$",
        outputs,
        flags=re.MULTILINE,
    )

    assert readme_match, f"riga XGBoost non trovata in {readme_path.name}"
    assert notebook_match, "riga XGBoost non trovata negli output del notebook"

    readme_values = np.array([float(x) for x in readme_match.groups()[:6]])
    notebook_values = np.array([float(x) for x in notebook_match.groups()[:6]])
    assert round(readme_values[0], 3) == notebook_values[0]
    np.testing.assert_array_equal(readme_values[1:], notebook_values[1:])
    assert readme_match.groups()[6:] == notebook_match.groups()[6:]


def test_notebook_non_contiene_claim_o_baseline_obsoleti():
    sources, _ = notebook_text()

    assert "45.5%/91.4%" not in sources
    assert "va davvero in produzione" not in sources
    assert "precision 0,490, recall 0,837" in sources
    assert "non documenta un deployment aziendale reale" in sources


def test_excel_to_prioritized_csv_in_a_new_process(tmp_path):
    for path, expected in EXPECTED_HASHES.items():
        assert sha256(path) == expected

    output = tmp_path / "retention_list.csv"
    environment = os.environ.copy()
    environment["PYTHONUTF8"] = "1"
    environment["PYTHONIOENCODING"] = "utf-8"
    process = subprocess.run(
        [
            sys.executable,
            str(SCRIPT),
            "--input", str(INPUT),
            "--model", str(MODEL),
            "--output", str(output),
        ],
        cwd=tmp_path,
        capture_output=True,
        text=True,
        encoding="utf-8",
        env=environment,
        check=False,
    )

    assert process.returncode == 0, (
        f"STDOUT:\n{process.stdout}\nSTDERR:\n{process.stderr}"
    )
    assert "Clienti analizzati  : 8" in process.stdout
    assert "Predizioni sopra soglia: 5" in process.stdout
    assert output.is_file()

    result = pd.read_csv(output, index_col="CustomerID")
    assert result.index.tolist() == EXPECTED_IDS
    np.testing.assert_allclose(
        result["proba"].to_numpy(), EXPECTED_PROBABILITIES, rtol=0, atol=1e-12
    )
    np.testing.assert_array_equal(result["decisione"].to_numpy(), np.ones(5, dtype=int))
    np.testing.assert_array_equal(result["decile"].to_numpy(), EXPECTED_DECILES)
    assert not result["motivi"].isna().any()
    for customer_id, markers in EXPECTED_REASON_MARKERS.items():
        motivo = result.loc[customer_id, "motivi"]
        assert all(marker in motivo for marker in markers)


@pytest.mark.parametrize(
    "colonna",
    ["Tenure Months", "Monthly Charges", "Total Charges"],
)
def test_cli_rifiuta_booleano_excel_nelle_colonne_numeriche(tmp_path, colonna):
    input_non_valido = tmp_path / f"booleano_{colonna.replace(' ', '_')}.xlsx"
    output = tmp_path / "risultato.csv"

    workbook = load_workbook(INPUT)
    foglio = workbook.active
    intestazioni = {
        cella.value: cella.column for cella in foglio[1]
    }
    foglio.cell(row=2, column=intestazioni[colonna], value=True)
    workbook.save(input_non_valido)
    workbook.close()

    controllo = load_workbook(input_non_valido, read_only=True, data_only=False)
    cella = controllo.active.cell(row=2, column=intestazioni[colonna])
    assert cella.value is True
    assert cella.data_type == "b"
    controllo.close()

    environment = os.environ.copy()
    environment["PYTHONUTF8"] = "1"
    environment["PYTHONIOENCODING"] = "utf-8"
    process = subprocess.run(
        [
            sys.executable,
            str(SCRIPT),
            "--input", str(input_non_valido),
            "--model", str(MODEL),
            "--output", str(output),
        ],
        cwd=tmp_path,
        capture_output=True,
        text=True,
        encoding="utf-8",
        env=environment,
        check=False,
    )

    assert process.returncode != 0
    assert colonna in process.stderr
    assert "riga Excel 2" in process.stderr
    assert not output.exists()


@pytest.mark.parametrize(
    "case,match",
    [
        ("missing_column", "colonna CustomerID"),
        ("null_id", "valori mancanti"),
        ("duplicate_id", "univoco"),
        ("whitespace_id", "vuoto o contenere soltanto spazi"),
        ("leading_space", "spazi iniziali o finali"),
        ("trailing_space", "spazi iniziali o finali"),
    ],
)
def test_customer_id_non_valido_viene_rifiutato(tmp_path, case, match):
    clienti = pd.read_excel(INPUT)

    if case == "missing_column":
        clienti = clienti.drop(columns=["CustomerID"])
    elif case == "null_id":
        clienti.loc[0, "CustomerID"] = None
    elif case == "duplicate_id":
        clienti.loc[1, "CustomerID"] = clienti.loc[0, "CustomerID"]
    elif case == "whitespace_id":
        clienti.loc[0, "CustomerID"] = "   "
    elif case == "leading_space":
        clienti.loc[0, "CustomerID"] = " DEMO-001"
    elif case == "trailing_space":
        clienti.loc[0, "CustomerID"] = "DEMO-001 "

    input_non_valido = tmp_path / f"{case}.xlsx"
    output = tmp_path / f"{case}.csv"
    clienti.to_excel(input_non_valido, index=False)

    with pytest.raises(ValueError, match=match):
        esegui(input_non_valido, MODEL, output)

    assert not output.exists()
