"""Esegue il ciclo mensile di inferenza da Excel a CSV."""

from __future__ import annotations

import argparse
from pathlib import Path

import joblib
from openpyxl import load_workbook
import pandas as pd

from churn_inference import predici_churn


CARTELLA_ESEMPIO = Path(__file__).resolve().parent
COLONNE_NUMERICHE_EXCEL = (
    "Tenure Months",
    "Monthly Charges",
    "Total Charges",
)


def valida_tipi_numerici_excel(input_excel: Path):
    """Rifiuta celle booleane prima che pandas le converta in 0/1."""
    workbook = load_workbook(
        input_excel,
        read_only=True,
        data_only=False,
        keep_links=False,
    )
    try:
        foglio = workbook.worksheets[0]
        posizioni = {}
        for cella in foglio[1]:
            if cella.value in COLONNE_NUMERICHE_EXCEL:
                posizioni.setdefault(cella.value, []).append(cella.column)

        booleani = {}
        for colonna, indici_colonna in posizioni.items():
            righe = []
            for indice_colonna in indici_colonna:
                for (cella,) in foglio.iter_rows(
                    min_row=2,
                    min_col=indice_colonna,
                    max_col=indice_colonna,
                ):
                    if cella.data_type == "b" or isinstance(cella.value, bool):
                        righe.append(cella.row)
            if righe:
                booleani[colonna] = sorted(set(righe))[:5]
    finally:
        workbook.close()

    if booleani:
        def descrivi_righe(righe):
            if len(righe) == 1:
                return f"riga Excel {righe[0]}"
            return f"righe Excel {righe}"

        dettagli = "; ".join(
            f"{colonna}: {descrivi_righe(righe)}"
            for colonna, righe in booleani.items()
        )
        raise ValueError(
            "Le colonne numeriche non accettano celle booleane; "
            f"valori non validi per colonna: {dettagli}."
        )


def leggi_argomenti(argv=None):
    parser = argparse.ArgumentParser(
        description="Calcola i punteggi di churn ed esporta le righe sopra soglia."
    )
    parser.add_argument(
        "--input",
        type=Path,
        default=CARTELLA_ESEMPIO / "examples" / "customers_sample.xlsx",
        help="File Excel con i clienti da analizzare.",
    )
    parser.add_argument(
        "--model",
        type=Path,
        default=CARTELLA_ESEMPIO / "pipeline_churn.joblib",
        help="Pipeline serializzata da utilizzare.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=CARTELLA_ESEMPIO / "outputs" / "churn_scores_above_threshold.csv",
        help="Percorso del CSV da generare.",
    )
    return parser.parse_args(argv)


def esegui(input_excel: Path, modello_joblib: Path, output_csv: Path):
    """Carica input e modello, calcola i punteggi e salva le righe sopra soglia."""
    valida_tipi_numerici_excel(input_excel)
    clienti_del_mese = pd.read_excel(input_excel)
    if "CustomerID" not in clienti_del_mese.columns:
        raise ValueError("Il file Excel deve contenere la colonna CustomerID.")
    if clienti_del_mese["CustomerID"].isna().any():
        raise ValueError("CustomerID non può contenere valori mancanti.")

    customer_id = clienti_del_mese["CustomerID"].astype("string")
    customer_id_normalizzato = customer_id.str.strip()
    if customer_id_normalizzato.eq("").any():
        raise ValueError("CustomerID non può essere vuoto o contenere soltanto spazi.")
    if customer_id.ne(customer_id_normalizzato).any():
        raise ValueError("CustomerID non può contenere spazi iniziali o finali.")
    if customer_id_normalizzato.duplicated().any():
        raise ValueError("CustomerID deve essere univoco nel file Excel.")

    clienti_del_mese = clienti_del_mese.set_index("CustomerID")
    pipeline = joblib.load(modello_joblib)
    risultati = predici_churn(
        pipeline,
        clienti_del_mese,
        con_motivazioni=True,
    ).sort_values(["decile", "proba"], ascending=[True, False])

    righe_sopra_soglia = risultati.loc[risultati["decisione"] == 1]
    output_csv.parent.mkdir(parents=True, exist_ok=True)
    righe_sopra_soglia.to_csv(output_csv)
    return risultati, righe_sopra_soglia


def main(argv=None):
    args = leggi_argomenti(argv)
    risultati, righe_sopra_soglia = esegui(args.input, args.model, args.output)

    print(f"Clienti analizzati  : {len(risultati)}")
    print(f"Predizioni sopra soglia: {len(righe_sopra_soglia)} (decisione = 1)\n")
    for cliente, riga in righe_sopra_soglia.iterrows():
        print(f"{cliente}  decile {riga['decile']}  p={riga['proba']:.2f}")
        print(f"    {riga['motivi']}")
    print(f"\nTabella esportata in {args.output.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
